from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "AI-Powered-Study-Buddy-Presentation.pptx"

NAVY = RGBColor(15, 31, 56)
INK = RGBColor(32, 43, 56)
MUTED = RGBColor(93, 108, 125)
BG = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 145, 150)
CORAL = RGBColor(235, 105, 91)
GOLD = RGBColor(237, 181, 71)
PALE_TEAL = RGBColor(225, 246, 245)
PALE_CORAL = RGBColor(253, 235, 232)
PALE_GOLD = RGBColor(255, 246, 220)
FONT = "Aptos"


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(0.7)
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def bullet_list(slide, items, x, y, w, h, size=16, color=INK, gap=0.08):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(size * gap)
        p.bullet = True
    return tb


def add_notes(slide, notes):
    try:
        tf = slide.notes_slide.notes_text_frame
        tf.text = notes
    except Exception:
        pass


def base(title, kicker="PROJECT DEFENSE"):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    box(slide, 0, 0, 13.333, 0.16, TEAL, TEAL, False)
    text(slide, kicker, 0.55, 0.38, 4.5, 0.25, 9, TEAL, True)
    text(slide, title, 0.55, 0.72, 11.9, 0.55, 27, NAVY, True)
    text(slide, f"{len(prs.slides):02d}", 12.25, 7.08, 0.5, 0.2, 9, MUTED, True, PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, accent=TEAL, fill=WHITE, body_size=14):
    box(slide, x, y, w, h, fill, RGBColor(224, 230, 237))
    box(slide, x, y, 0.08, h, accent, accent, False)
    text(slide, heading, x + 0.22, y + 0.18, w - 0.4, 0.3, 15, NAVY, True)
    text(slide, body, x + 0.22, y + 0.62, w - 0.4, h - 0.76, body_size, INK)


def node(slide, x, y, w, h, label, fill=WHITE, accent=TEAL, size=13):
    box(slide, x, y, w, h, fill, accent)
    text(slide, label, x + 0.08, y + 0.08, w - 0.16, h - 0.16, size, NAVY, True, PP_ALIGN.CENTER)


def arrow(slide, x1, y1, x2, y2, color=TEAL):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.8)
    line.line.end_arrowhead = True


def section_slide(title, subtitle, number):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    box(slide, 0, 0, 13.333, 0.16, CORAL, CORAL, False)
    text(slide, f"0{number}", 0.7, 1.2, 1.4, 0.6, 28, GOLD, True)
    text(slide, title, 0.7, 2.15, 10.8, 0.8, 38, WHITE, True)
    text(slide, subtitle, 0.75, 3.25, 9.8, 0.7, 18, RGBColor(206, 220, 232))
    add_notes(slide, f"Transition into the {title.lower()} section. {subtitle}")
    return slide


# 1 Cover
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
box(slide, 0, 0, 13.333, 0.2, CORAL, CORAL, False)
box(slide, 9.3, 0.2, 4.03, 7.3, TEAL, TEAL, False)
text(slide, "ACADEMIC PROJECT PRESENTATION", 0.75, 0.72, 6.0, 0.3, 11, GOLD, True)
text(slide, "AI-Powered\nStudy Buddy", 0.72, 1.55, 7.6, 1.55, 38, WHITE, True)
text(slide, "A full-stack learning assistant for document-grounded study", 0.78, 3.55, 7.1, 0.55, 19, RGBColor(206, 220, 232))
text(slide, "Student: [Information Required]\nRoll / Registration No.: [Information Required]\nCourse, Department, Institute: [Information Required]\nGuide / Supervisor: [Information Required]\nAcademic year / submission date: [Information Required]", 0.8, 5.0, 6.7, 1.35, 14, WHITE)
node(slide, 9.8, 1.2, 2.7, 0.75, "UPLOAD", PALE_GOLD, GOLD)
node(slide, 9.8, 2.35, 2.7, 0.75, "UNDERSTAND", PALE_TEAL, TEAL)
node(slide, 9.8, 3.5, 2.7, 0.75, "PRACTICE", PALE_CORAL, CORAL)
node(slide, 9.8, 4.65, 2.7, 0.75, "IMPROVE", WHITE, GOLD)
arrow(slide, 11.15, 1.95, 11.15, 2.3, WHITE)
arrow(slide, 11.15, 3.1, 11.15, 3.45, WHITE)
arrow(slide, 11.15, 4.25, 11.15, 4.6, WHITE)
add_notes(slide, "Introduce the project as an AI-assisted study workflow. Replace all bracketed academic metadata before submission. The repository does not provide student, institution, guide, or submission information.")

# 2 Overview
slide = base("Project Overview")
text(slide, "Study Buddy turns a student's own study material into a connected learning loop.", 0.6, 1.45, 8.2, 0.5, 20, NAVY, True)
node(slide, 0.8, 2.45, 2.25, 1.0, "STUDY\nMATERIAL", PALE_GOLD, GOLD)
node(slide, 3.55, 2.45, 2.25, 1.0, "AI\nPROCESSING", PALE_TEAL, TEAL)
node(slide, 6.3, 2.45, 2.25, 1.0, "ACTIVE\nPRACTICE", PALE_CORAL, CORAL)
node(slide, 9.05, 2.45, 2.25, 1.0, "PROGRESS\nFEEDBACK", WHITE, TEAL)
for x in (3.05, 5.8, 8.55): arrow(slide, x, 2.95, x + 0.45, 2.95)
card(slide, 0.8, 4.25, 3.55, 1.55, "Inputs", "PDF, DOCX, TXT, or Markdown files; optional document context for tutor questions.", GOLD, PALE_GOLD)
card(slide, 4.85, 4.25, 3.55, 1.55, "Core capabilities", "Summaries, quizzes, flashcards, streaming AI tutor, and document management.", TEAL, PALE_TEAL)
card(slide, 8.9, 4.25, 3.55, 1.55, "Audience", "Students who want one workspace for turning notes into study activities.", CORAL, PALE_CORAL)
add_notes(slide, "Explain the product in one sentence: users upload study content, AI transforms it into study artifacts, and activity is recorded for progress views. The target audience is inferred from the UI and README wording; validate the intended academic segment if needed.")

# 3 Background
slide = base("Background and Motivation")
card(slide, 0.75, 1.55, 3.75, 3.95, "The study friction", "Students often move between notes, a summarizer, a quiz tool, flashcards, and a progress tracker. That breaks context and makes revision effort harder to organize.", CORAL, PALE_CORAL, 16)
card(slide, 4.8, 1.55, 3.75, 3.95, "The project response", "Study Buddy keeps source material and the resulting learning activities connected through a single authenticated application.", TEAL, PALE_TEAL, 16)
card(slide, 8.85, 1.55, 3.75, 3.95, "Why it matters", "The workflow supports both understanding and retrieval practice: explain a concept, test recall, schedule review, and inspect progress.", GOLD, PALE_GOLD, 16)
text(slide, "Motivation supported by repository product description; no external user research or statistics supplied.", 0.8, 6.05, 11.5, 0.3, 11, MUTED)
add_notes(slide, "Frame motivation as a workflow problem, not a quantified market claim. Do not add adoption or learning-outcome statistics because the repository contains none.")

# 4 Problem
slide = base("Problem Statement")
node(slide, 0.8, 2.0, 2.7, 1.2, "FRAGMENTED\nSTUDY TOOLS", PALE_CORAL, CORAL)
node(slide, 5.3, 2.0, 2.7, 1.2, "LOW-CONTEXT\nREVISION", PALE_GOLD, GOLD)
node(slide, 9.8, 2.0, 2.7, 1.2, "WEAK\nFEEDBACK LOOP", PALE_TEAL, TEAL)
arrow(slide, 3.55, 2.6, 5.2, 2.6, CORAL)
arrow(slide, 8.05, 2.6, 9.7, 2.6, GOLD)
card(slide, 0.8, 4.15, 3.5, 1.25, "Problem", "Study content and study actions are disconnected.", CORAL, WHITE, 15)
card(slide, 4.9, 4.15, 3.5, 1.25, "Impact", "Students spend effort preparing material instead of practicing it.", GOLD, WHITE, 15)
card(slide, 9.0, 4.15, 3.5, 1.25, "Need", "A unified, content-aware study assistant with feedback.", TEAL, WHITE, 15)
text(slide, "Project-specific framing derived from implemented modules and README overview.", 0.8, 6.1, 10, 0.3, 11, MUTED)
add_notes(slide, "State the problem clearly and avoid claiming that the system has proven to improve grades. The implemented response is integration: source documents feed summaries, quizzes, flashcards, and tutor context.")

# 5 Existing vs proposed
slide = base("Existing Approach and Project Gap")
text(slide, "The repository does not name a specific competitor or baseline system. This comparison describes the workflow gap, not measured superiority.", 0.65, 1.35, 11.9, 0.45, 13, MUTED)
card(slide, 0.8, 2.1, 5.4, 3.35, "Typical fragmented approach", "- Notes stored separately\n- Manual summary and question creation\n- Practice history split across tools\n- Little connection between source content and feedback", CORAL, PALE_CORAL, 16)
card(slide, 7.1, 2.1, 5.4, 3.35, "Study Buddy approach", "- Upload and process source material\n- Generate summaries, quizzes, and flashcards\n- Ask a document-aware tutor\n- Record study activity and topic progress", TEAL, PALE_TEAL, 16)
text(slide, "[Information Required] Add literature review or named existing systems if required by the academic format.", 0.8, 6.05, 11.5, 0.3, 11, CORAL, True)
add_notes(slide, "Be precise: the codebase supports the proposed workflow, but it does not provide a formal comparative evaluation or named existing-system study. Add those only from your report or research sources.")

# 6 Proposed system
slide = base("Proposed System")
text(slide, "One authenticated workspace connects content ingestion, AI generation, practice, and analytics.", 0.7, 1.35, 11.5, 0.4, 19, NAVY, True)
node(slide, 0.75, 2.3, 2.1, 1.0, "UPLOAD\nDOCUMENT", PALE_GOLD, GOLD)
node(slide, 3.25, 2.3, 2.1, 1.0, "EXTRACT\nTEXT", PALE_TEAL, TEAL)
node(slide, 5.75, 2.3, 2.1, 1.0, "GENERATE\nLEARNING", PALE_CORAL, CORAL)
node(slide, 8.25, 2.3, 2.1, 1.0, "PRACTICE\nAND CHAT", WHITE, TEAL)
node(slide, 10.75, 2.3, 1.8, 1.0, "TRACK\nPROGRESS", PALE_GOLD, GOLD)
for x in (2.85, 5.35, 7.85, 10.35): arrow(slide, x, 2.8, x + 0.35, 2.8)
card(slide, 1.0, 4.35, 3.3, 1.35, "Content-aware", "Summaries, quizzes, flashcards, and tutor prompts use extracted document text.", TEAL, WHITE, 14)
card(slide, 5.0, 4.35, 3.3, 1.35, "Adaptive review", "SM-2 updates intervals and next review dates from quality ratings.", CORAL, WHITE, 14)
card(slide, 9.0, 4.35, 3.3, 1.35, "Observable activity", "Quiz attempts, reviews, sessions, topics, and mastery feed analytics.", GOLD, WHITE, 14)
add_notes(slide, "Walk left to right. The key design idea is continuity: the uploaded material is the common context for AI-generated activities, while user activity is persisted for feedback.")

# 7 Objectives
slide = base("Project Objectives")
objectives = [
    "Accept and process common study-document formats.",
    "Generate summaries in five supported formats.",
    "Generate configurable multiple-choice quizzes with explanations.",
    "Generate flashcards and schedule reviews with SM-2.",
    "Provide six AI tutor modes, including optional document context.",
    "Record learning activity and expose topic-level analytics.",
]
for i, item in enumerate(objectives):
    x = 0.85 if i < 3 else 6.75
    y = 1.55 + (i % 3) * 1.45
    box(slide, x, y, 5.6, 1.0, WHITE, RGBColor(224, 230, 237))
    text(slide, f"0{i + 1}", x + 0.2, y + 0.22, 0.5, 0.35, 16, CORAL if i % 2 else TEAL, True)
    text(slide, item, x + 0.85, y + 0.17, 4.45, 0.55, 15, INK, True)
text(slide, "Objectives are phrased from implemented routes, services, models, and frontend screens.", 0.85, 6.15, 11, 0.3, 11, MUTED)
add_notes(slide, "These are implementation-aligned objectives. They are not outcome metrics. Explain that achievement should be demonstrated through feature walkthroughs and tests, while learning impact requires a separate evaluation.")

# 8 Requirements
slide = base("Scope and Requirements")
card(slide, 0.75, 1.5, 3.85, 4.65, "Functional requirements", "- Clerk sign-in and local profile sync\n- Upload PDF, DOCX, TXT, Markdown\n- Background extraction and status tracking\n- Generate summaries, quizzes, flashcards\n- Tutor chat with six modes and SSE\n- Review cards and submit quiz attempts\n- Persist activity and analytics", TEAL, PALE_TEAL, 14)
card(slide, 4.75, 1.5, 3.85, 4.65, "Non-functional requirements", "- Async API and database I/O\n- User-scoped resource access\n- File type and size validation\n- Parameterized ORM queries\n- CORS allowlist\n- Rate limiting\n- Standardized error responses\n- Light/dark UI persistence", GOLD, PALE_GOLD, 14)
card(slide, 8.75, 1.5, 3.85, 4.65, "Scope boundary", "Included: local full-stack workflow, AI integration, relational persistence, browser UI, Docker Compose setup.\n\n[Information Required] Confirm deployment target, supported browsers, accessibility criteria, and formal performance targets.", CORAL, PALE_CORAL, 14)
add_notes(slide, "Distinguish what the code implements from requirements that should be specified in the report. The repository describes local Docker development, not a verified production deployment.")

# 9 Stack
slide = base("Technology Stack")
stack = [
    ("Frontend", "HTML5 + Tailwind CSS CDN + Vanilla JS ES6 modules", "No build step; direct browser workflow", TEAL),
    ("Backend", "Python 3.12+ + FastAPI + Uvicorn", "Async REST API and background tasks", CORAL),
    ("Data", "MySQL 8.0 + SQLAlchemy async + aiomysql + Alembic", "Relational persistence and migrations", GOLD),
    ("AI", "OpenAI SDK + configurable OpenAI-compatible base URL", "Streaming tutor and JSON generation", TEAL),
    ("Identity", "Clerk JS SDK + RS256 JWT/JWKS verification", "External identity with user isolation", CORAL),
    ("Operations", "Docker Compose + SlowAPI + Pydantic v2", "Local services, limits, validation/config", GOLD),
]
for i, (head, tech, why, accent) in enumerate(stack):
    y = 1.4 + i * 0.82
    box(slide, 0.8, y, 2.0, 0.62, accent, accent)
    text(slide, head, 0.95, y + 0.15, 1.7, 0.25, 13, WHITE, True, PP_ALIGN.CENTER)
    text(slide, tech, 3.1, y + 0.08, 5.25, 0.4, 14, NAVY, True)
    text(slide, why, 8.65, y + 0.08, 3.7, 0.4, 13, MUTED)
text(slide, "The repository also lists pypdf, python-docx, python-jose, httpx, aiofiles, and python-multipart.", 0.85, 6.55, 11.6, 0.3, 11, MUTED)
add_notes(slide, "Explain choices in terms of the actual code: FastAPI and async SQLAlchemy support the API/data path, Clerk handles identity, the OpenAI SDK supports configurable providers, and vanilla JS keeps frontend setup lightweight.")

# 10 Architecture
slide = base("System Architecture")
node(slide, 0.65, 2.45, 2.15, 1.0, "USER\nBROWSER", PALE_GOLD, GOLD)
node(slide, 3.45, 1.75, 2.25, 1.0, "HTML + JS\nTAILWIND", PALE_TEAL, TEAL)
node(slide, 3.45, 3.2, 2.25, 1.0, "CLERK\nAUTH", PALE_CORAL, CORAL)
node(slide, 6.45, 2.45, 2.25, 1.0, "FASTAPI\nROUTERS", WHITE, TEAL)
node(slide, 9.45, 1.75, 2.25, 1.0, "MYSQL\nDATA", PALE_GOLD, GOLD)
node(slide, 9.45, 3.2, 2.25, 1.0, "OPENAI-\nCOMPATIBLE LLM", PALE_CORAL, CORAL)
arrow(slide, 2.8, 2.95, 3.35, 2.25)
arrow(slide, 2.8, 2.95, 3.35, 3.7, CORAL)
arrow(slide, 5.75, 2.25, 6.35, 2.8)
arrow(slide, 5.75, 3.7, 6.35, 3.1, CORAL)
arrow(slide, 8.8, 2.8, 9.35, 2.25, GOLD)
arrow(slide, 8.8, 3.1, 9.35, 3.7, CORAL)
text(slide, "REST + SSE", 4.4, 1.25, 1.4, 0.25, 11, TEAL, True)
text(slide, "Async ORM", 8.8, 1.25, 1.4, 0.25, 11, GOLD, True)
text(slide, "Streaming AI", 8.75, 4.55, 1.7, 0.25, 11, CORAL, True)
text(slide, "Architecture follows docs/architecture.md and app/main.py router registration.", 0.8, 6.25, 11.5, 0.3, 11, MUTED)
add_notes(slide, "Describe the two external dependencies explicitly: Clerk for identity and an OpenAI-compatible LLM provider for generation. The application layer connects the browser to MySQL and the AI layer through REST and SSE.")

# 11 Workflow
slide = base("End-to-End Workflow")
steps = [("1", "Sign in", GOLD), ("2", "Upload", TEAL), ("3", "Extract", CORAL), ("4", "Generate", TEAL), ("5", "Practice", GOLD), ("6", "Analyze", CORAL)]
for i, (num, label, accent) in enumerate(steps):
    x = 0.7 + i * 2.1
    node(slide, x, 2.25, 1.55, 1.1, f"{num}\n{label}", WHITE, accent, 14)
    if i < len(steps) - 1: arrow(slide, x + 1.58, 2.8, x + 2.0, 2.8)
text(slide, "Upload path", 0.8, 4.25, 1.4, 0.3, 14, NAVY, True)
text(slide, "Validate MIME and size -> save UUID filename -> create uploaded record -> BackgroundTask -> extract text -> ready or failed", 2.2, 4.23, 10.2, 0.45, 15, INK)
text(slide, "Learning path", 0.8, 5.05, 1.4, 0.3, 14, NAVY, True)
text(slide, "Ready document -> call LLM -> persist summary / quiz / cards -> review or submit -> update progress and study session", 2.2, 5.03, 10.2, 0.45, 15, INK)
add_notes(slide, "Emphasize the asynchronous boundary: the upload response is not blocked by extraction. AI generation only proceeds for documents whose processing status is ready. Activity logging and learning progress are persisted alongside study actions.")

# 12 Modules
slide = base("Functional Modules")
modules = [
    ("Documents", "Upload, list, open, delete; background extraction", GOLD),
    ("Summaries", "Five summary types generated from extracted text", TEAL),
    ("Quizzes", "Generate, take, score, explain, record attempt", CORAL),
    ("Flashcards", "Generate, flip, rate, schedule with SM-2", GOLD),
    ("AI Tutor", "Six modes, optional document context, SSE stream", TEAL),
    ("Learning", "Topics, mastery, recommendations, sessions", CORAL),
    ("Analytics", "Counts, average score, streak, strong/weak topics", GOLD),
    ("Account", "Clerk auth, profile sync, settings, theme", TEAL),
]
for i, (head, body, accent) in enumerate(modules):
    x = 0.75 + (i % 4) * 3.1
    y = 1.55 + (i // 4) * 2.25
    card(slide, x, y, 2.65, 1.55, head, body, accent, WHITE, 12)
add_notes(slide, "Use this slide to orient the examiner before the detailed demonstration. Each module corresponds to a frontend page or backend router/service visible in the repository.")

# 13 Database
slide = base("Database Design")
node(slide, 0.75, 2.5, 1.55, 0.75, "USERS", PALE_GOLD, GOLD)
node(slide, 3.0, 1.45, 1.75, 0.75, "DOCUMENTS", PALE_TEAL, TEAL)
node(slide, 3.0, 3.55, 1.75, 0.75, "TOPICS", PALE_CORAL, CORAL)
node(slide, 5.8, 1.0, 1.75, 0.75, "SUMMARIES", WHITE, TEAL)
node(slide, 5.8, 2.0, 1.75, 0.75, "QUIZZES", WHITE, CORAL)
node(slide, 5.8, 3.0, 1.75, 0.75, "FLASHCARDS", WHITE, GOLD)
node(slide, 5.8, 4.0, 1.75, 0.75, "PROGRESS", WHITE, TEAL)
node(slide, 8.6, 1.5, 1.9, 0.75, "QUESTIONS", PALE_TEAL, TEAL)
node(slide, 8.6, 2.65, 1.9, 0.75, "ATTEMPTS", PALE_CORAL, CORAL)
node(slide, 8.6, 3.8, 1.9, 0.75, "ANSWERS", PALE_GOLD, GOLD)
node(slide, 11.0, 2.65, 1.75, 0.75, "SESSIONS", WHITE, CORAL)
for a, b in [((2.3, 2.85), (2.95, 1.82)), ((2.3, 2.85), (2.95, 3.92)), ((4.8, 1.82), (5.75, 1.38)), ((4.8, 1.82), (5.75, 2.38)), ((4.8, 3.92), (5.75, 3.38)), ((4.8, 3.92), (5.75, 4.38)), ((7.6, 2.38), (8.55, 1.88)), ((7.6, 2.38), (8.55, 3.03)), ((10.55, 3.03), (10.95, 3.03))]: arrow(slide, *a, *b)
text(slide, "Initial Alembic migration defines 11 tables: users, documents, topics, summaries, flashcards, quizzes, quiz_questions, quiz_attempts, quiz_answers, learning_progress, study_sessions.", 0.8, 6.1, 11.8, 0.45, 11, MUTED)
add_notes(slide, "The migration defines 11 tables, although the README architecture sketch says 10 tables. Present the migration as the source of truth and flag the README count discrepancy for correction before submission. Primary keys are string UUID-style IDs; foreign keys connect users and learning artifacts.")

# 14 Implementation
slide = base("Detailed Implementation")
card(slide, 0.75, 1.45, 3.8, 4.5, "Document pipeline", "1. Validate MIME type and size\n2. Sanitize original filename\n3. Store with UUID-based path\n4. Create uploaded DB record\n5. Run extraction in BackgroundTasks\n6. Mark ready or failed", GOLD, PALE_GOLD, 15)
card(slide, 4.8, 1.45, 3.8, 4.5, "Learning logic", "Quiz score = correct / total * 100\n\nMastery update:\nnew = 0.7 * old + 0.3 * quiz score\n\nFlashcard quality 0-5 changes repetitions, ease factor, interval, and next_review_at.", TEAL, PALE_TEAL, 15)
card(slide, 8.85, 1.45, 3.8, 4.5, "AI interaction", "Prompts are stored as text files. Extracted text is truncated before LLM calls. Tutor responses stream as text/event-stream; summaries and quizzes are persisted after generation.", CORAL, PALE_CORAL, 15)
add_notes(slide, "This is the technical core. Explain the three control points: safe ingestion, deterministic progress updates, and prompt-driven AI generation. The formulas are directly represented in learning_service.py and flashcards.py.")

# 15 UI
slide = base("User Interface Evidence")
text(slide, "The repository contains functional HTML screens, but no captured screenshots were supplied.", 0.75, 1.4, 11.3, 0.4, 20, NAVY, True)
for i, (label, desc, accent) in enumerate([
    ("Sign in", "Clerk sign-in mount", CORAL), ("Documents", "Upload and processing status", GOLD),
    ("AI Tutor", "Mode selector and streaming chat", TEAL), ("Flashcards", "Flip and quality rating", CORAL),
    ("Quizzes", "Question flow and result panel", GOLD), ("Analytics", "Stats, topics, streak, activity", TEAL),
]):
    x = 0.8 + (i % 3) * 4.1
    y = 2.2 + (i // 3) * 1.55
    box(slide, x, y, 3.55, 1.1, WHITE, RGBColor(224, 230, 237))
    box(slide, x + 0.16, y + 0.18, 0.58, 0.7, accent, accent)
    text(slide, "UI", x + 0.26, y + 0.39, 0.38, 0.2, 11, WHITE, True, PP_ALIGN.CENTER)
    text(slide, label, x + 0.95, y + 0.18, 2.25, 0.25, 15, NAVY, True)
    text(slide, desc, x + 0.95, y + 0.55, 2.35, 0.25, 12, MUTED)
text(slide, "[Screenshot Required] Replace the UI evidence cards with actual application screenshots before submission.", 0.8, 5.75, 11.5, 0.4, 16, CORAL, True)
add_notes(slide, "Show the actual running application here during the defense. The source files prove that these screens exist, but screenshots are evidence of rendered behavior and were not included in the workspace.")

# 16 Testing
slide = base("Testing and Validation")
text(slide, "The test suite focuses on pure logic and validation helpers; it does not constitute a full end-to-end application evaluation.", 0.75, 1.35, 11.7, 0.4, 15, MUTED)
headers = ["Area", "Verified test intent", "Evidence status"]
for x, h in zip([0.8, 3.0, 8.5], headers):
    box(slide, x, 2.0, [2.0, 5.3, 3.9][headers.index(h)], 0.55, NAVY, NAVY, False)
    text(slide, h, x + 0.12, 2.15, [1.7, 4.9, 3.5][headers.index(h)], 0.2, 12, WHITE, True)
rows = [
    ("Quiz scoring", "Perfect, zero, partial, rounding, zero-total", "Repository tests"),
    ("SM-2", "Intervals, reset, ease factor, clamp, future date", "Repository tests"),
    ("Mastery", "First score and weighted moving average", "Repository tests"),
    ("File validation", "Extensions, MIME types, sanitization", "Repository tests"),
    ("Text extraction", "Cleaning, text/Markdown, page estimate", "Repository tests"),
    ("API integration", "Routes with real DB/LLM/auth services", "[Information Required]"),
]
for i, row in enumerate(rows):
    y = 2.55 + i * 0.56
    fill = WHITE if i % 2 == 0 else RGBColor(239, 243, 247)
    for x, w, val in zip([0.8, 3.0, 8.5], [2.0, 5.3, 3.9], row):
        box(slide, x, y, w, 0.55, fill, RGBColor(224, 230, 237), False)
        text(slide, val, x + 0.12, y + 0.14, w - 0.22, 0.22, 11, CORAL if "Required" in val else INK, "Required" not in val)
text(slide, "README badge reports 73 passing tests; independently rerun and record the actual result before submission.", 0.8, 6.25, 11.5, 0.3, 11, GOLD, True)
add_notes(slide, "Describe the tests as unit-style and standalone. Do not claim integration, UI, performance, or security test completion unless you run and document those checks. The README badge is a claim that should be verified in the current environment.")

# 17 Results
slide = base("Results and Demonstrated Outcomes")
card(slide, 0.8, 1.55, 3.7, 3.9, "Implemented functional outcomes", "- Accepted file formats and processing states\n- AI generation endpoints for summaries, quizzes, cards\n- Streaming tutor response path\n- Quiz scoring and persisted attempts\n- SM-2 scheduling state\n- Analytics aggregation endpoints", TEAL, PALE_TEAL, 15)
card(slide, 4.85, 1.55, 3.7, 3.9, "Evidence currently available", "- Source implementation\n- Alembic schema\n- Architecture documentation\n- Repository test modules\n- Frontend page structures\n\nThese support design and code claims, not measured user outcomes.", GOLD, PALE_GOLD, 15)
card(slide, 8.9, 1.55, 3.7, 3.9, "Required before submission", "[Screenshot Required]\n[Measured performance Required]\n[Observed output examples Required]\n[Evaluation dataset / user study Required, if applicable]", CORAL, PALE_CORAL, 15)
text(slide, "No accuracy, latency, user-count, grade-improvement, or production metrics are present in the supplied project materials.", 0.8, 6.05, 11.5, 0.3, 12, CORAL, True)
add_notes(slide, "This slide is intentionally honest. Demonstrate working behavior live or add captured evidence. Do not fill the empty evidence categories with invented metrics.")

# 18 Security
slide = base("Security and Reliability")
card(slide, 0.75, 1.45, 5.75, 4.7, "Implemented in code", "- Clerk JWT verification using RS256 and JWKS\n- User ownership checks on resources\n- MIME allowlist and 50 MB configured limit\n- UUID storage filenames and filename sanitization\n- SQLAlchemy ORM parameterized queries\n- Frontend HTML escaping helper\n- CORS allowlist and SlowAPI rate limiting\n- Standardized exception responses", TEAL, PALE_TEAL, 14)
card(slide, 6.8, 1.45, 5.75, 4.7, "Recommended verification / improvement", "- Add integration tests for authorization across every route\n- Confirm production HTTPS and secret management\n- Verify upload content sniffing beyond client MIME claims\n- Add durable job queue for scale beyond process-local BackgroundTasks\n- Define backup, recovery, observability, and availability targets\n- Validate prompt/data privacy policy with real provider", CORAL, PALE_CORAL, 14)
add_notes(slide, "Separate implementation from recommendation. The repository documents these controls, but production security posture and operational reliability are not demonstrated by the supplied files alone.")

# 19 Limitations
slide = base("Limitations")
limitations = [
    ("Evidence gap", "No screenshots, measured performance, or evaluation dataset supplied.", CORAL),
    ("AI dependency", "Generation quality depends on the configured external or local LLM provider.", TEAL),
    ("Processing model", "FastAPI BackgroundTasks are process-local; durable job orchestration is not shown.", GOLD),
    ("Scope", "Repository documents local Docker development; production deployment is not verified.", CORAL),
    ("Analytics", "Progress is activity-based and no learning-effectiveness study is included.", TEAL),
    ("Documentation mismatch", "README sketch says 10 tables; initial migration defines 11.", GOLD),
]
for i, (head, body, accent) in enumerate(limitations):
    x = 0.8 + (i % 2) * 6.05
    y = 1.5 + (i // 2) * 1.55
    card(slide, x, y, 5.45, 1.12, head, body, accent, WHITE, 13)
add_notes(slide, "These limitations are grounded in the repository. The table-count discrepancy is a documentation issue, not necessarily a runtime defect, and should be reconciled before the viva.")

# 20 Future
slide = base("Future Scope")
card(slide, 0.8, 1.5, 3.7, 4.45, "Short-term", "- Add route-level integration tests\n- Capture UI and output evidence\n- Reconcile schema documentation\n- Add richer profile data validation\n- Improve document error reporting", TEAL, PALE_TEAL, 15)
card(slide, 4.85, 1.5, 3.7, 4.45, "Medium-term", "- Durable background job queue\n- Conversation history persistence\n- Better topic extraction and tagging\n- Expanded analytics visualizations\n- Export of summaries and study plans", GOLD, PALE_GOLD, 15)
card(slide, 8.9, 1.5, 3.7, 4.45, "Long-term", "- Production deployment and observability\n- Provider-level privacy controls\n- Multi-device experience\n- Evaluation of learning outcomes\n- Scale and reliability engineering", CORAL, PALE_CORAL, 15)
add_notes(slide, "Keep future work connected to current architecture. Prioritize evidence, durable processing, evaluation, and operations before adding unrelated technologies.")

# 21 Contribution
slide = base("Project Contribution and Learning")
text(slide, "[Information Required] Replace this slide with the student's verified individual contribution and learning reflection.", 0.8, 1.45, 11.6, 0.45, 17, CORAL, True)
card(slide, 0.8, 2.35, 3.65, 2.8, "Technical areas evidenced", "Async API design\nRelational schema and migrations\nBrowser-to-API integration\nLLM prompt integration\nTesting and input validation", TEAL, PALE_TEAL, 15)
card(slide, 4.85, 2.35, 3.65, 2.8, "Engineering practices", "Separation of routes and services\nUser-scoped resources\nStructured error responses\nConfiguration through environment variables\nDocker-based local setup", GOLD, PALE_GOLD, 15)
card(slide, 8.9, 2.35, 3.65, 2.8, "Add before submission", "Student role / module ownership\nChallenges encountered\nDecisions defended\nSkills learned\nTeam contribution split", CORAL, PALE_CORAL, 15)
add_notes(slide, "Do not present repository-wide implementation as individual contribution unless it is true. Replace the required fields with the student's own documented work.")

# 22 Conclusion
slide = base("Conclusion")
text(slide, "Study Buddy implements a connected study workflow around a student's own material.", 0.8, 1.45, 11.5, 0.5, 24, NAVY, True)
for i, (head, body, accent) in enumerate([
    ("Problem addressed", "Fragmented preparation and practice workflow", CORAL),
    ("System delivered", "Authenticated full-stack learning assistant", TEAL),
    ("Technical contribution", "Async API, relational persistence, LLM integration, SM-2 review", GOLD),
    ("Evidence position", "Core implementation is present; screenshots and outcome measurements remain required", CORAL),
]):
    y = 2.45 + i * 0.82
    box(slide, 1.0, y, 11.2, 0.62, WHITE, RGBColor(224, 230, 237))
    text(slide, head, 1.25, y + 0.17, 2.3, 0.22, 13, accent, True)
    text(slide, body, 3.8, y + 0.15, 7.9, 0.25, 15, INK, True)
add_notes(slide, "Close by answering the examiner's four questions: what problem, what was built, what technical design, and what evidence remains to be added. Avoid claiming learning improvement without an evaluation.")

# 23 References
slide = base("References and Evidence Sources")
refs = [
    "Project README.md - product overview, features, stack, setup, and reported test badge.",
    "docs/architecture.md - architecture, request flows, security model, and design decisions.",
    "backend/alembic/versions/001_initial_migration.py - database tables, keys, and relationships.",
    "backend/app/main.py and backend/app/api/routes/ - API registration and implemented route behavior.",
    "backend/app/services/learning_service.py - mastery calculations and learning progress behavior.",
    "backend/app/api/routes/flashcards.py - SM-2 scheduling implementation.",
    "backend/app/core/security.py and backend/app/api/deps.py - Clerk token verification and ownership dependency.",
    "tests/ - standalone validation, extraction, scoring, SM-2, mastery, and security-oriented checks.",
]
bullet_list(slide, refs, 0.9, 1.45, 11.4, 4.9, 15, INK, 0.12)
text(slide, "[Information Required] Add any external papers, official documentation, datasets, or APIs actually cited in the final report.", 0.9, 6.35, 11.4, 0.3, 12, CORAL, True)
add_notes(slide, "These references are repository evidence sources. Add formal external references only if they were actually used. Do not list a paper, dataset, or website merely because it is related to the topic.")

# 24 Viva
slide = base("Viva Preparation: Likely Questions")
card(slide, 0.75, 1.4, 3.85, 4.9, "Basic", "Q: What is the project?\nA: An AI-powered learning assistant that turns uploaded study material into summaries, quizzes, flashcards, and tutor interactions.\n\nQ: Why this topic?\nA: To connect content preparation, active recall, and progress tracking in one workflow.", TEAL, PALE_TEAL, 13)
card(slide, 4.75, 1.4, 3.85, 4.9, "Technical", "Q: Why FastAPI?\nA: It fits the async REST and streaming API design.\n\nQ: Why Clerk?\nA: It provides signed identity tokens while the backend enforces local ownership.\n\nQ: How does SM-2 work?\nA: Quality ratings update repetitions, ease factor, interval, and next review date.", GOLD, PALE_GOLD, 13)
card(slide, 8.75, 1.4, 3.85, 4.9, "Critical", "Q: What are limitations?\nA: No measured outcome study, production deployment evidence, or durable job queue is supplied.\n\nQ: How would you improve it?\nA: Add integration evidence, durable processing, privacy controls, observability, and learning evaluation.", CORAL, PALE_CORAL, 13)
add_notes(slide, "Answer in terms of the implementation. For questions about student contribution, performance, screenshots, deployment, or evaluation, use the required information markers until those facts are documented.")

# 25 Q&A
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
box(slide, 0, 0, 13.333, 0.2, CORAL, CORAL, False)
text(slide, "THANK YOU", 0.8, 1.65, 11.7, 0.7, 42, WHITE, True, PP_ALIGN.CENTER)
text(slide, "Questions and discussion", 0.8, 2.65, 11.7, 0.45, 22, RGBColor(206, 220, 232), False, PP_ALIGN.CENTER)
text(slide, "AI-Powered Study Buddy\n[Student name and contact: Information Required]", 0.8, 4.75, 11.7, 0.7, 16, GOLD, True, PP_ALIGN.CENTER)
add_notes(slide, "Invite questions. Keep the live demonstration or repository open for follow-up questions about routes, schema, security, and algorithms.")


prs.core_properties.title = "AI-Powered Study Buddy - Academic Project Presentation"
prs.core_properties.subject = "Project-aligned technical presentation"
prs.core_properties.author = "Generated from repository evidence"
prs.save(OUT)
print(f"Wrote {OUT}")
print(f"Slides: {len(prs.slides)}")